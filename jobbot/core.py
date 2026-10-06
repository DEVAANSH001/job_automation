import hashlib
import html
import json
import os
import re
import sqlite3
import threading
from datetime import datetime, timezone
from pathlib import Path

import httpx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

LOCK = threading.Lock()


def now():
    return datetime.now(timezone.utc).isoformat()


def config():
    return json.loads(Path(os.getenv('CONFIG_PATH', 'config.json')).read_text(encoding='utf-8'))


def root():
    path = Path(os.getenv('DATA_DIR', 'data'))
    path.mkdir(parents=True, exist_ok=True)
    return path


def db():
    conn = sqlite3.connect(root() / 'jobs.db', timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA journal_mode=WAL')
    conn.execute('CREATE TABLE IF NOT EXISTS jobs (id TEXT PRIMARY KEY, source TEXT, company TEXT, title TEXT, location TEXT, url TEXT, description TEXT, score INTEGER, status TEXT, resume TEXT, discovered TEXT, updated TEXT, applied TEXT, note TEXT)')
    conn.execute('CREATE TABLE IF NOT EXISTS events (time TEXT, job_id TEXT, action TEXT, detail TEXT)')
    conn.commit()
    return conn


def event(job_id, action, detail=''):
    with db() as conn:
        conn.execute('INSERT INTO events VALUES (?,?,?,?)', (now(), job_id, action, detail))


def text(value):
    return html.unescape(re.sub('<[^>]+>', ' ', str(value or '')))


def collect(board):
    source, slug = board['source'], board['slug']
    if not re.fullmatch(r'[\w-]+', slug):
        raise ValueError('Invalid company board slug')
    company = board.get('company', slug)
    with httpx.Client(timeout=30, follow_redirects=False) as client:
        if source == 'greenhouse':
            response = client.get(f'https://boards-api.greenhouse.io/v1/boards/{slug}/jobs', params={'content': 'true'})
        elif source == 'lever':
            response = client.get(f'https://api.lever.co/v0/postings/{slug}', params={'mode': 'json'})
        elif source == 'ashby':
            response = client.get(f'https://api.ashbyhq.com/posting-api/job-board/{slug}')
        else:
            raise ValueError('Supported sources: greenhouse, lever, ashby')
        response.raise_for_status()
        payload = response.json()
    jobs = payload if source == 'lever' else payload['jobs']
    for job in jobs:
        location = job.get('location', '')
        if isinstance(location, dict):
            location = location.get('name', '')
        if source == 'lever':
            location = job.get('categories', {}).get('location', '')
        yield {
            'id': hashlib.sha256(f'{source}:{slug}:{job["id"]}'.encode()).hexdigest()[:24],
            'source': source, 'company': company, 'title': job.get('title', job.get('text', '')),
            'location': location,
            'url': job.get('applyUrl') or job.get('absolute_url') or job.get('hostedUrl') or job.get('jobUrl'),
            'description': text(job.get('content') or job.get('descriptionPlain') or job.get('descriptionHtml') or job.get('description', ''))
        }


def score(job, cfg):
    title = job['title'].lower()
    if any(x.lower() in title for x in cfg.get('exclude_titles', [])):
        return 0
    roles = cfg.get('roles', [])
    locations = cfg.get('locations', [])
    if roles and not any(x.lower() in title for x in roles):
        return 0
    if locations and not any(x.lower() in job['location'].lower() for x in locations):
        return 0
    max_years = cfg.get('max_required_years')
    if max_years is not None:
        description = job['description'].lower()
        patterns = [
            r'(\d+)\s*(?:\+|[-\u2013\u2014\ufffd]\s*\d+)?\s*years?\s+of\s+(?:professional\s+|industry\s+|commercial\s+)?(?:software\s+)?(?:engineering\s+|development\s+)?experience',
            r'(\d+)\s*(?:\+|[-\u2013\u2014\ufffd]\s*\d+)?\s*years?\s+(?:of\s+)?experience',
            r'(?:professional\s+|industry\s+|commercial\s+)?(?:software\s+)?(?:engineering\s+|development\s+)?experience.{0,35}?(\d+)\s*(?:\+|[-\u2013\u2014\ufffd]\s*\d+)?\s*years?',
        ]
        required = [int(value) for pattern in patterns for value in re.findall(pattern, description)]
        if required and min(required) > int(max_years):
            return 0
    skills = cfg['profile'].get('skills', [])
    matched = sum(x.lower() in job['description'].lower() for x in skills)
    return 50 + round(50 * matched / len(skills)) if skills else 50


def update(job_id, **fields):
    allowed = {'status', 'resume', 'applied', 'note'}
    if not fields.keys() <= allowed:
        raise ValueError('Invalid update')
    fields['updated'] = now()
    with db() as conn:
        conn.execute('UPDATE jobs SET ' + ','.join(f'{k}=?' for k in fields) + ' WHERE id=?', [*fields.values(), job_id])


def rows():
    with db() as conn:
        return [dict(row) for row in conn.execute('SELECT * FROM jobs ORDER BY score DESC, discovered DESC')]


def excel():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Jobs'
    headers = ['id', 'company', 'title', 'location', 'source', 'score', 'status', 'url', 'resume', 'discovered', 'updated', 'applied', 'note']
    sheet.append(headers)
    for job in rows():
        sheet.append([job[k] for k in headers])
    audit = workbook.create_sheet('Activity')
    audit.append(['UTC time', 'Job ID', 'Action', 'Detail'])
    with db() as conn:
        for row in conn.execute('SELECT * FROM events ORDER BY time DESC'):
            audit.append(list(row))
    for tab in workbook:
        tab.freeze_panes = 'A2'
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243B53')
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or '')) for c in column) + 2))
        for row in tab.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Job descriptions must never become Excel formulas.
    temporary = root() / 'jobs.tmp.xlsx'
    workbook.save(temporary)
    temporary.replace(root() / 'jobs.xlsx')


def prepare(job, cfg):
    profile = cfg['profile']
    if not profile.get('name') or not profile.get('email') or not profile.get('experience'):
        update(job['id'], status='needs_profile', note='Complete name, email and experience in config.json')
        return
    styles = getSampleStyleSheet()
    content = []
    def paragraph(value, style='Normal'):
        content.append(Paragraph(html.escape(str(value)), styles[style]))
        content.append(Spacer(1, 8))
    paragraph(profile['name'], 'Title')
    paragraph(' | '.join(str(profile.get(k, '')) for k in ['email', 'phone', 'location']))
    for key in ['linkedin', 'github', 'portfolio']:
        if profile.get(key):
            paragraph(profile[key])
    if profile.get('summary'):
        paragraph(profile['summary'])
    skills = profile.get('skills', [])
    ordered = sorted(skills, key=lambda s: s.lower() in job['description'].lower(), reverse=True)
    paragraph('Skills', 'Heading2')
    paragraph(', '.join(ordered))
    paragraph('Experience', 'Heading2')
    for experience in profile['experience']:
        paragraph(experience['title'] + ' | ' + experience['company'] + ' | ' + experience.get('dates', ''), 'Heading3')
        bullets = sorted(experience.get('bullets', []), key=lambda b: sum(s.lower() in b.lower() and s.lower() in job['description'].lower() for s in skills), reverse=True)
        for bullet in bullets:
            paragraph('• ' + bullet)
    paragraph('Education', 'Heading2')
    for education in profile.get('education', []):
        paragraph(education)
    if profile.get('projects'):
        paragraph('Projects', 'Heading2')
        for project in profile['projects']:
            paragraph(project['name'], 'Heading3')
            if project.get('url'):
                paragraph(project['url'])
            for bullet in project.get('bullets', []):
                paragraph('- ' + bullet)
    if profile.get('achievements'):
        paragraph('Achievements', 'Heading2')
        for achievement in profile['achievements']:
            paragraph('- ' + achievement)
    path = root() / f'resume-{job["id"]}.pdf'
    SimpleDocTemplate(str(path)).build(content)
    update(job['id'], status='prepared', resume=str(path), note='Verified profile content reordered for JD relevance')
    event(job['id'], 'resume_prepared')


def run():
    if not LOCK.acquire(blocking=False):
        return {'status': 'already_running'}
    try:
        cfg = config()
        for board in cfg.get('boards', []):
            try:
                for job in collect(board):
                    rating = score(job, cfg)
                    timestamp = now()
                    with db() as conn:
                        conn.execute('INSERT INTO jobs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET title=excluded.title,location=excluded.location,url=excluded.url,description=excluded.description,score=excluded.score,updated=excluded.updated',
                                     (*[job[k] for k in ['id','source','company','title','location','url','description']], rating, 'discovered', '', timestamp, timestamp, '', ''))
                event('', 'collection_ok', f'{board["source"]}:{board["slug"]}')
            except Exception as exc:
                event('', 'collection_error', f'{board.get("slug")}: {type(exc).__name__}')
        current_jobs = rows()
        rescored = [(score(job, cfg), job['id']) for job in current_jobs]
        with db() as conn:
            conn.executemany('UPDATE jobs SET score=? WHERE id=?', rescored)
        for job, (rating, _) in zip(current_jobs, rescored):
            job['score'] = rating
            if job['score'] == 0 and job['status'] in ('discovered', 'prepared', 'needs_profile', 'filtered'):
                update(job['id'], status='filtered', note='Outside configured title, location, or experience rules')
                continue
            if job['score'] > 0 and job['status'] == 'filtered':
                update(job['id'], status='discovered', note='')
                job['status'] = 'discovered'
            if job['status'] in ('discovered', 'needs_profile') and job['score'] >= cfg.get('minimum_score', 50) and job['score'] > 0:
                try:
                    prepare(job, cfg)
                except Exception as exc:
                    update(job['id'], status='error', note=type(exc).__name__)
                    event(job['id'], 'prepare_error', type(exc).__name__)
        if cfg.get('auto_submit'):
            from jobbot.apply import apply
            for job in rows():
                if job['status'] == 'prepared' and job['score'] >= cfg.get('minimum_score', 50):
                    apply(job, cfg)
        excel()
        return {'status': 'completed', 'jobs': len(rows())}
    finally:
        LOCK.release()
